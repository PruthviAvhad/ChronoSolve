"""Excel import and template export.

A coordinator does not want to type 156 session rows. They think in *subjects*:
"Data Structures, for SE-A, taught by Prof. Mehta, three one-hour classes a
week." So the Subjects sheet carries weekly requirements and this module expands
them into the individual sessions the solver schedules.

Workbook layout (sheet names are matched case-insensitively):

    Rooms           id | name | capacity | type
    Faculty         id | name | max_daily | max_weekly | max_consecutive
    Batches         id | name | strength
    Subjects        code | name | batch | faculty | sessions_per_week |
                    hours_per_session | room_type | elective_group | headcount
    Unavailability  kind | id | day | start | end          (optional)

Nothing is imported until the whole workbook validates. A partially applied
import would leave a coordinator with a half-replaced department and no way to
tell which half, so errors are collected and reported together.
"""

from __future__ import annotations

import io
import math
from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from ..domain.models import (
    DAYS,
    PERIOD_END,
    PERIOD_START,
    Batch,
    Calendar,
    Faculty,
    Instance,
    Room,
    RoomType,
    Session,
)

ERROR = "error"
WARNING = "warning"


@dataclass(frozen=True, slots=True)
class ImportIssue:
    severity: str
    sheet: str
    row: int | None
    message: str

    def render(self) -> str:
        where = self.sheet + (f" row {self.row}" if self.row else "")
        return f"[{self.severity}] {where}: {self.message}"


@dataclass(slots=True)
class ImportReport:
    issues: list[ImportIssue] = field(default_factory=list)
    counts: dict[str, int] = field(default_factory=dict)

    @property
    def errors(self) -> list[ImportIssue]:
        return [i for i in self.issues if i.severity == ERROR]

    @property
    def warnings(self) -> list[ImportIssue]:
        return [i for i in self.issues if i.severity == WARNING]

    @property
    def ok(self) -> bool:
        return not self.errors

    def add(self, severity: str, sheet: str, row: int | None, message: str) -> None:
        self.issues.append(ImportIssue(severity, sheet, row, message))

    def render(self) -> str:
        lines = []
        if self.counts:
            lines.append(
                "  read " + ", ".join(f"{v} {k}" for k, v in self.counts.items())
            )
        for issue in self.issues:
            lines.append(f"  {issue.render()}")
        if self.ok and not self.issues:
            lines.append("  no problems found")
        return "\n".join(lines)


# ----------------------------------------------------------------------
# reading helpers
# ----------------------------------------------------------------------


def _rows(ws) -> list[tuple[int, dict]]:
    """(row_number, {column: value}) for every populated row below the header."""
    it = ws.iter_rows(values_only=True)
    try:
        raw_header = next(it)
    except StopIteration:
        return []

    header = [str(h).strip().lower() if h is not None else "" for h in raw_header]
    out: list[tuple[int, dict]] = []
    for number, values in enumerate(it, start=2):
        if all(v is None or str(v).strip() == "" for v in values):
            continue
        out.append((number, dict(zip(header, values))))
    return out


def _text(row: dict, *names: str) -> str:
    for name in names:
        value = row.get(name)
        if value is not None and str(value).strip():
            return str(value).strip()
    return ""


def _int(row: dict, *names: str) -> int | None:
    """Parsed integer, None when absent, -1 when present but unreadable."""
    for name in names:
        value = row.get(name)
        if value is None or str(value).strip() == "":
            continue
        try:
            return int(float(str(value).strip()))
        except ValueError:
            return -1
    return None


def _time(value) -> str | None:
    """Accept "14:00", "2 PM", or an Excel time cell."""
    if value is None or str(value).strip() == "":
        return None
    if hasattr(value, "hour"):
        return f"{value.hour:02d}:{value.minute:02d}"

    text = str(value).strip().upper()
    suffix = ""
    for token in ("AM", "PM"):
        if token in text:
            suffix = token
            text = text.replace(token, "").strip()
    parts = text.split(":")
    try:
        hour = int(parts[0])
        minute = int(parts[1]) if len(parts) > 1 else 0
    except ValueError:
        return None
    if suffix == "PM" and hour != 12:
        hour += 12
    if suffix == "AM" and hour == 12:
        hour = 0
    if not (0 <= hour <= 23 and 0 <= minute <= 59):
        return None
    return f"{hour:02d}:{minute:02d}"


def _day(value) -> int | None:
    if value is None:
        return None
    text = str(value).strip().lower()
    for index, name in enumerate(DAYS):
        if text.startswith(name.lower()):
            return index
    if text.isdigit() and 0 <= int(text) < len(DAYS):
        return int(text)
    return None


# ----------------------------------------------------------------------
# import
# ----------------------------------------------------------------------


def read_workbook(
    data: bytes, name: str = "Imported department"
) -> tuple[Instance | None, ImportReport]:
    """Parse and validate a workbook. Returns (instance, report).

    The instance is None whenever the report holds errors, so callers should
    check `report.ok` rather than testing the instance.
    """
    report = ImportReport()
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # openpyxl raises several unrelated types
        report.add(ERROR, "workbook", None, f"could not be opened: {exc}")
        return None, report

    sheets = {n.strip().lower(): n for n in wb.sheetnames}
    for required in ("rooms", "faculty", "batches", "subjects"):
        if required not in sheets:
            report.add(
                ERROR,
                "workbook",
                None,
                f"missing required sheet '{required.title()}' "
                f"(found: {', '.join(wb.sheetnames)})",
            )
    if not report.ok:
        return None, report

    calendar = Calendar()
    teaching = {s.id for s in calendar.teaching_slots}

    # --- rooms --------------------------------------------------------
    rooms: dict[str, Room] = {}
    for number, row in _rows(wb[sheets["rooms"]]):
        rid = _text(row, "id", "room", "room_id")
        if not rid:
            report.add(ERROR, "Rooms", number, "missing id")
            continue
        if rid in rooms:
            report.add(ERROR, "Rooms", number, f"duplicate room id {rid!r}")
            continue
        capacity = _int(row, "capacity", "seats")
        if capacity is None or capacity <= 0:
            report.add(
                ERROR, "Rooms", number, f"{rid}: capacity must be a positive number"
            )
            continue
        kind = _text(row, "type", "room_type").upper() or "LECTURE"
        if kind not in ("LECTURE", "LAB"):
            report.add(
                ERROR,
                "Rooms",
                number,
                f"{rid}: type must be LECTURE or LAB, got {kind!r}",
            )
            continue
        rooms[rid] = Room(
            id=rid,
            name=_text(row, "name") or rid,
            capacity=capacity,
            room_type=RoomType(kind),
        )

    # --- faculty ------------------------------------------------------
    people: dict[str, Faculty] = {}
    for number, row in _rows(wb[sheets["faculty"]]):
        fid = _text(row, "id", "faculty_id", "code")
        if not fid:
            report.add(ERROR, "Faculty", number, "missing id")
            continue
        if fid in people:
            report.add(ERROR, "Faculty", number, f"duplicate faculty id {fid!r}")
            continue
        people[fid] = Faculty(
            id=fid,
            name=_text(row, "name") or fid,
            max_daily_load=_int(row, "max_daily", "max_daily_load") or 5,
            max_weekly_load=_int(row, "max_weekly", "max_weekly_load") or 18,
            max_consecutive=_int(row, "max_consecutive") or 3,
        )

    # --- batches ------------------------------------------------------
    batches: dict[str, Batch] = {}
    for number, row in _rows(wb[sheets["batches"]]):
        bid = _text(row, "id", "batch", "batch_id")
        if not bid:
            report.add(ERROR, "Batches", number, "missing id")
            continue
        if bid in batches:
            report.add(ERROR, "Batches", number, f"duplicate batch id {bid!r}")
            continue
        strength = _int(row, "strength", "students", "size")
        if strength is None or strength <= 0:
            report.add(
                ERROR, "Batches", number, f"{bid}: strength must be a positive number"
            )
            continue
        batches[bid] = Batch(id=bid, name=_text(row, "name") or bid, strength=strength)

    # --- unavailability ----------------------------------------------
    blocked_faculty: dict[str, set[int]] = {}
    blocked_rooms: dict[str, set[int]] = {}
    blocked_batches: dict[str, set[int]] = {}
    if "unavailability" in sheets:
        for number, row in _rows(wb[sheets["unavailability"]]):
            kind = _text(row, "kind", "type").lower()
            target = _text(row, "id", "target", "who")
            day = _day(row.get("day"))
            start = _time(row.get("start")) or PERIOD_START[0]
            end = _time(row.get("end")) or "23:59"

            if kind not in ("faculty", "room", "batch"):
                report.add(
                    ERROR,
                    "Unavailability",
                    number,
                    f"kind must be faculty, room or batch, got {kind!r}",
                )
                continue
            if day is None:
                report.add(
                    ERROR,
                    "Unavailability",
                    number,
                    f"day must be one of {', '.join(DAYS)}",
                )
                continue
            pool = {"faculty": people, "room": rooms, "batch": batches}[kind]
            if target not in pool:
                report.add(ERROR, "Unavailability", number, f"unknown {kind} {target!r}")
                continue
            if start >= end:
                report.add(
                    ERROR,
                    "Unavailability",
                    number,
                    f"{target}: {start}-{end} ends before it starts",
                )
                continue

            slots = {
                sl.id
                for sl in calendar.slots_on(day)
                if not sl.is_lunch and start <= PERIOD_START[sl.period] < end
            }
            if not slots:
                # Silently ignoring this would hide a typo -- "20:00" for
                # "2:00 PM" would produce a rule that quietly does nothing.
                report.add(
                    ERROR,
                    "Unavailability",
                    number,
                    f"{target}: {start}-{end} covers no teaching period "
                    f"({PERIOD_START[0]}-{PERIOD_END[-1]}, lunch excluded)",
                )
                continue
            bucket = {
                "faculty": blocked_faculty,
                "room": blocked_rooms,
                "batch": blocked_batches,
            }[kind]
            bucket.setdefault(target, set()).update(slots)

    people = {
        fid: Faculty(
            id=f.id,
            name=f.name,
            unavailable=frozenset(blocked_faculty.get(fid, set())),
            max_daily_load=f.max_daily_load,
            max_weekly_load=f.max_weekly_load,
            max_consecutive=f.max_consecutive,
        )
        for fid, f in people.items()
    }
    rooms = {
        rid: Room(
            id=r.id,
            name=r.name,
            capacity=r.capacity,
            room_type=r.room_type,
            unavailable=frozenset(blocked_rooms.get(rid, set())),
        )
        for rid, r in rooms.items()
    }
    batches = {
        bid: Batch(
            id=b.id,
            name=b.name,
            strength=b.strength,
            unavailable=frozenset(blocked_batches.get(bid, set())),
        )
        for bid, b in batches.items()
    }

    # --- subjects -> sessions ----------------------------------------
    sessions: list[Session] = []
    seen: set[tuple[str, str]] = set()
    for number, row in _rows(wb[sheets["subjects"]]):
        code = _text(row, "code", "subject_code")
        batch_id = _text(row, "batch", "batch_id", "division")
        faculty_id = _text(row, "faculty", "faculty_id", "teacher")

        if not code:
            report.add(ERROR, "Subjects", number, "missing subject code")
            continue
        if batch_id not in batches:
            report.add(ERROR, "Subjects", number, f"{code}: unknown batch {batch_id!r}")
            continue
        if faculty_id not in people:
            report.add(
                ERROR, "Subjects", number, f"{code}: unknown faculty {faculty_id!r}"
            )
            continue
        if (code, batch_id) in seen:
            report.add(
                ERROR, "Subjects", number, f"{code} is listed twice for {batch_id}"
            )
            continue
        seen.add((code, batch_id))

        per_week = _int(row, "sessions_per_week", "sessions", "per_week") or 1
        hours = _int(row, "hours_per_session", "hours", "duration") or 1
        kind = _text(row, "room_type", "type").upper() or "LECTURE"
        group = _text(row, "elective_group", "elective") or None
        headcount = _int(row, "headcount", "cohort")

        if per_week <= 0 or hours <= 0:
            report.add(
                ERROR,
                "Subjects",
                number,
                f"{code}: sessions_per_week and hours_per_session must be positive",
            )
            continue
        if kind not in ("LECTURE", "LAB"):
            report.add(
                ERROR, "Subjects", number, f"{code}: room_type must be LECTURE or LAB"
            )
            continue
        if hours > calendar.periods:
            report.add(
                ERROR,
                "Subjects",
                number,
                f"{code}: a {hours}-hour block cannot fit in a "
                f"{calendar.periods}-period day",
            )
            continue

        room_type = RoomType(kind)
        seats = headcount if headcount and headcount > 0 else batches[batch_id].strength
        if not any(
            r.room_type is room_type and r.capacity >= seats for r in rooms.values()
        ):
            report.add(
                ERROR,
                "Subjects",
                number,
                f"{code}: no {kind.lower()} room seats {seats} students",
            )
            continue

        for k in range(per_week):
            sessions.append(
                Session(
                    id=f"{batch_id}-{code}-{k + 1}",
                    subject_code=code,
                    subject_name=_text(row, "name", "subject", "subject_name") or code,
                    batch_id=batch_id,
                    faculty_id=faculty_id,
                    duration=hours,
                    room_type=room_type,
                    elective_group=f"{group}-{k + 1}" if group else None,
                    headcount=headcount if headcount and headcount > 0 else None,
                )
            )

    report.counts = {
        "rooms": len(rooms),
        "faculty": len(people),
        "batches": len(batches),
        "sessions": len(sessions),
    }

    # --- whole-workbook sanity ---------------------------------------
    if not sessions and report.ok:
        report.add(ERROR, "Subjects", None, "no sessions were produced")

    for fid, f in people.items():
        load = sum(s.duration for s in sessions if s.faculty_id == fid)
        free = len(teaching) - len(f.unavailable)
        if load > f.max_weekly_load:
            report.add(
                ERROR,
                "Faculty",
                None,
                f"{f.name} is assigned {load}h but their weekly cap is "
                f"{f.max_weekly_load}h",
            )
        elif load > free:
            report.add(
                ERROR,
                "Faculty",
                None,
                f"{f.name} is assigned {load}h but is free for only {free} periods",
            )
        elif load and free - load <= 1:
            report.add(
                WARNING,
                "Faculty",
                None,
                f"{f.name} has almost no free periods left ({load}h of {free})",
            )

    for bid, b in batches.items():
        hours = sum(
            s.duration
            for s in sessions
            if s.batch_id == bid
            and (not s.elective_group or s.elective_group.endswith("-1"))
        )
        if hours > len(teaching):
            report.add(
                ERROR,
                "Subjects",
                None,
                f"{b.id} needs {hours} periods but the week has only {len(teaching)}",
            )

    idle = [f.name for fid, f in people.items() if not any(s.faculty_id == fid for s in sessions)]
    if idle:
        report.add(
            WARNING,
            "Faculty",
            None,
            f"{len(idle)} faculty teach nothing: {', '.join(sorted(idle)[:5])}",
        )

    if not report.ok:
        return None, report

    return (
        Instance(
            name=name,
            calendar=calendar,
            rooms=list(rooms.values()),
            faculty=list(people.values()),
            batches=list(batches.values()),
            sessions=sessions,
        ),
        report,
    )


# ----------------------------------------------------------------------
# template export
# ----------------------------------------------------------------------


def to_workbook(instance: Instance) -> bytes:
    """The current department written back out in the import format.

    Downloading this, editing it and re-importing is the intended round trip,
    which also keeps the template in step with the importer by construction.
    """
    wb = Workbook()

    ws = wb.active
    ws.title = "Rooms"
    ws.append(["id", "name", "capacity", "type"])
    for r in instance.rooms:
        ws.append([r.id, r.name, r.capacity, r.room_type.value])

    ws = wb.create_sheet("Faculty")
    ws.append(["id", "name", "max_daily", "max_weekly", "max_consecutive"])
    for f in instance.faculty:
        ws.append([f.id, f.name, f.max_daily_load, f.max_weekly_load, f.max_consecutive])

    ws = wb.create_sheet("Batches")
    ws.append(["id", "name", "strength"])
    for b in instance.batches:
        ws.append([b.id, b.name, b.strength])

    ws = wb.create_sheet("Subjects")
    ws.append(
        [
            "code",
            "name",
            "batch",
            "faculty",
            "sessions_per_week",
            "hours_per_session",
            "room_type",
            "elective_group",
            "headcount",
        ]
    )
    # Collapse sessions back into the weekly requirement they came from.
    grouped: dict[tuple[str, str], list[Session]] = {}
    for s in instance.sessions:
        grouped.setdefault((s.batch_id, s.subject_code), []).append(s)
    for (batch_id, code), group in grouped.items():
        first = group[0]
        elective = (
            first.elective_group.rsplit("-", 1)[0] if first.elective_group else None
        )
        ws.append(
            [
                code,
                first.subject_name,
                batch_id,
                first.faculty_id,
                len(group),
                first.duration,
                first.room_type.value,
                elective,
                first.headcount,
            ]
        )

    ws = wb.create_sheet("Unavailability")
    ws.append(["kind", "id", "day", "start", "end"])
    for f in instance.faculty:
        for day, start, end in _windows(instance.calendar, f.unavailable):
            ws.append(["faculty", f.id, DAYS[day], start, end])
    for r in instance.rooms:
        for day, start, end in _windows(instance.calendar, r.unavailable):
            ws.append(["room", r.id, DAYS[day], start, end])
    for b in instance.batches:
        for day, start, end in _windows(instance.calendar, b.unavailable):
            ws.append(["batch", b.id, DAYS[day], start, end])

    for sheet in wb.worksheets:
        for column in range(1, sheet.max_column + 1):
            letter = get_column_letter(column)
            longest = max(
                (len(str(c.value)) for c in sheet[letter] if c.value is not None),
                default=8,
            )
            sheet.column_dimensions[letter].width = min(max(longest + 2, 10), 34)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _windows(calendar: Calendar, slots: frozenset[int]) -> list[tuple[int, str, str]]:
    """Collapse blocked slot ids into contiguous day/start/end windows."""
    out: list[tuple[int, str, str]] = []
    for day in range(calendar.days):
        periods = sorted(sl.period for sl in calendar.slots_on(day) if sl.id in slots)
        if not periods:
            continue
        run_start = previous = periods[0]
        for period in periods[1:] + [None]:
            if period is not None and period == previous + 1:
                previous = period
                continue
            out.append((day, PERIOD_START[run_start], PERIOD_END[previous]))
            if period is not None:
                run_start = previous = period
    return out
