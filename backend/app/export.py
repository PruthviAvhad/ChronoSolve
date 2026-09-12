"""Excel, PDF and calendar export.

Each function returns bytes (or text) rather than writing to disk, so the same
code serves an HTTP download, the console demo, and any future batch job.

The exports are deliberately not just a dump of the grid: they carry the solver
status and the independent validation result alongside the schedule, so a
printed timetable can still be traced back to the run that produced it.
"""

from __future__ import annotations

import io
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from .domain.models import DAYS, PERIOD_END, PERIOD_START, Instance, Timetable
from .solver.metrics import ScheduleDiff, schedule_metrics
from .solver.validate import validate

HEADER_FILL = PatternFill("solid", fgColor="1F2C45")
LUNCH_FILL = PatternFill("solid", fgColor="F2F2F2")
LAB_FILL = PatternFill("solid", fgColor="DDEBF7")
CHANGED_FILL = PatternFill("solid", fgColor="FFF2CC")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _grid_map(
    instance: Instance, timetable: Timetable, keep
) -> dict[tuple[int, int], list]:
    """(day, period) -> sessions visible in this view, continuations included."""
    cal = instance.calendar
    out: dict[tuple[int, int], list] = {}
    for s in instance.sessions:
        p = timetable.placements.get(s.id)
        if p is None or not keep(s, p):
            continue
        start = cal.by_id[p.timeslot_id]
        for k in range(s.duration):
            slot = cal.by_id.get(start.id + k)
            if slot is None:
                continue
            out.setdefault((slot.day, slot.period), []).append((s, p, k > 0))
    return out


# ----------------------------------------------------------------------
# Excel
# ----------------------------------------------------------------------


def to_xlsx(
    instance: Instance,
    timetable: Timetable,
    diff: ScheduleDiff | None = None,
    label: str = "published",
) -> bytes:
    """Workbook: summary, a grid per division, flat session list, changes."""
    cal = instance.calendar
    changed = {c.session_id for c in diff.changes} if diff else set()
    report = validate(instance, timetable)
    quality = schedule_metrics(instance, timetable)

    wb = Workbook()

    # --- summary -----------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 46

    def row(k: str, v, bold: bool = False) -> None:
        r = ws.max_row + 1
        ws.cell(r, 1, k).font = Font(bold=True, color="1F2C45")
        cell = ws.cell(r, 2, v)
        if bold:
            cell.font = Font(bold=True)

    ws.cell(1, 1, "ChronoSolve - Adaptive Constraint-Optimized Scheduling").font = Font(
        bold=True, size=14, color="1F2C45"
    )
    ws.cell(2, 1, instance.name).font = Font(italic=True, color="595959")
    ws.append([])

    row("Timetable", label)
    row("Exported", datetime.now().strftime("%Y-%m-%d %H:%M"))
    row("Solver", "Google OR-Tools CP-SAT")
    row("Solver status", timetable.status, bold=True)
    if timetable.objective is not None:
        row("Objective", timetable.objective)
        row("Best bound", timetable.best_bound)
    row("Solve time (s)", round(timetable.solve_seconds, 2))
    ws.append([])

    row("Sessions", len(instance.sessions))
    row("Contact hours", instance.total_contact_hours)
    row(
        "Divisions / faculty / rooms",
        f"{len(instance.batches)} / {len(instance.faculty)} / {len(instance.rooms)}",
    )
    ws.append([])

    ws.cell(ws.max_row + 1, 1, "Hard constraints (verified independently)").font = Font(
        bold=True, size=12, color="1F2C45"
    )
    for key, count in report.counts.items():
        row(key.replace("_", " "), count)
    ws.append([])

    ws.cell(ws.max_row + 1, 1, "Schedule quality").font = Font(
        bold=True, size=12, color="1F2C45"
    )
    row("Student idle hours", quality.student_idle_hours)
    row("Faculty idle hours", quality.faculty_idle_hours)
    row("Room utilisation %", round(quality.room_utilisation_pct, 1))
    row("Faculty load spread (h)", quality.faculty_load_spread)

    if diff:
        ws.append([])
        ws.cell(ws.max_row + 1, 1, "Repair").font = Font(
            bold=True, size=12, color="1F2C45"
        )
        row("Total sessions", diff.total)
        row("Unchanged", diff.unchanged)
        row("Changed", diff.changed)
        row("Schedule retention %", round(diff.retention_pct, 1), bold=True)

    # --- one grid per division ---------------------------------------
    for batch in instance.batches:
        sheet = wb.create_sheet(batch.id[:31])
        sheet.cell(
            1, 1, f"{batch.id} - {batch.name} ({batch.strength} students)"
        ).font = Font(bold=True, size=12, color="1F2C45")

        sheet.cell(3, 1, "Time").font = Font(bold=True, color="FFFFFF")
        sheet.cell(3, 1).fill = HEADER_FILL
        for d in range(cal.days):
            c = sheet.cell(3, d + 2, DAYS[d])
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HEADER_FILL
            c.alignment = Alignment(horizontal="center")

        grid = _grid_map(instance, timetable, lambda s, p, b=batch: s.batch_id == b.id)
        stacked_per_row: list[int] = []
        for period in range(cal.periods):
            r = period + 4
            stacked_per_row.append(
                max((len(grid.get((d, period), [])) for d in range(cal.days)), default=1)
            )
            sheet.cell(
                r, 1, f"{PERIOD_START[period]}-{PERIOD_END[period]}"
            ).font = Font(bold=True)
            sheet.cell(r, 1).border = BORDER
            for d in range(cal.days):
                cell = sheet.cell(r, d + 2)
                cell.border = BORDER
                cell.alignment = Alignment(
                    wrap_text=True, vertical="center", horizontal="center"
                )
                slot = cal.by_id[d * cal.periods + period]
                if slot.is_lunch:
                    cell.value = "LUNCH"
                    cell.fill = LUNCH_FILL
                    continue
                entries = grid.get((d, period), [])
                if not entries:
                    continue
                parts = []
                for s, p, cont in entries:
                    suffix = " (cont.)" if cont else ""
                    parts.append(
                        f"{s.subject_code}{suffix}\n{p.room_id} - "
                        f"{instance.faculty_by_id[s.faculty_id].name}"
                    )
                cell.value = "\n".join(parts)
                if any(s.id in changed for s, _, _ in entries):
                    cell.fill = CHANGED_FILL
                elif any(s.is_lab for s, _, _ in entries):
                    cell.fill = LAB_FILL

        sheet.column_dimensions["A"].width = 14
        for d in range(cal.days):
            sheet.column_dimensions[get_column_letter(d + 2)].width = 26
        for period in range(cal.periods):
            # Two text lines per stacked subject, ~15pt each.
            sheet.row_dimensions[period + 4].height = (
                32 * max(stacked_per_row[period], 1) + 2
            )

    # --- flat session list -------------------------------------------
    ws = wb.create_sheet("All sessions")
    headers = [
        "Session",
        "Subject code",
        "Subject",
        "Division",
        "Faculty",
        "Day",
        "Start",
        "End",
        "Hours",
        "Room",
        "Type",
        "Elective group",
        "Changed",
    ]
    ws.append(headers)
    for i in range(1, len(headers) + 1):
        ws.cell(1, i).font = Font(bold=True, color="FFFFFF")
        ws.cell(1, i).fill = HEADER_FILL

    for s in sorted(instance.sessions, key=lambda s: (s.batch_id, s.subject_code)):
        p = timetable.placements.get(s.id)
        if p is None:
            continue
        slot = cal.by_id[p.timeslot_id]
        ws.append(
            [
                s.id,
                s.subject_code,
                s.subject_name,
                s.batch_id,
                instance.faculty_by_id[s.faculty_id].name,
                DAYS[slot.day],
                PERIOD_START[slot.period],
                PERIOD_END[min(slot.period + s.duration - 1, cal.periods - 1)],
                s.duration,
                p.room_id,
                s.room_type.value,
                s.elective_group or "",
                "yes" if s.id in changed else "",
            ]
        )
    for i, width in enumerate([22, 13, 32, 10, 18, 7, 8, 8, 7, 9, 10, 22, 9], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # --- changes ------------------------------------------------------
    if diff and diff.changes:
        ws = wb.create_sheet("Changes")
        heads = ["Subject", "Division", "Faculty", "From", "To", "Kind"]
        ws.append(heads)
        for i in range(1, len(heads) + 1):
            ws.cell(1, i).font = Font(bold=True, color="FFFFFF")
            ws.cell(1, i).fill = HEADER_FILL
        for c in diff.changes:
            ws.append(
                [
                    c.subject_name,
                    c.batch_id,
                    c.faculty_name,
                    f"{c.from_label} {c.from_room}",
                    f"{c.to_label} {c.to_room}",
                    "time" if c.moved_time else "room",
                ]
            )
        for i, width in enumerate([32, 10, 18, 20, 20, 8], start=1):
            ws.column_dimensions[get_column_letter(i)].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ----------------------------------------------------------------------
# PDF
# ----------------------------------------------------------------------


def to_pdf(
    instance: Instance,
    timetable: Timetable,
    diff: ScheduleDiff | None = None,
    label: str = "published",
) -> bytes:
    """One landscape page per division, plus a verification page."""
    cal = instance.calendar
    changed = {c.session_id for c in diff.changes} if diff else set()
    report = validate(instance, timetable)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=f"ChronoSolve timetable ({label})",
    )
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle(
        "h1",
        parent=styles["Heading1"],
        fontSize=16,
        textColor=colors.HexColor("#1f2c45"),
    )
    sub = ParagraphStyle(
        "sub", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#666666")
    )
    cellstyle = ParagraphStyle("cell", parent=styles["Normal"], fontSize=7, leading=8.5)

    story: list = []

    for index, batch in enumerate(instance.batches):
        if index:
            story.append(PageBreak())
        story.append(Paragraph(f"{batch.id} - {batch.name}", h1))
        story.append(
            Paragraph(
                f"{instance.name} &middot; {batch.strength} students &middot; "
                f"timetable &ldquo;{label}&rdquo; &middot; solver {timetable.status} "
                f"&middot; {report.total} hard-constraint violations",
                sub,
            )
        )
        story.append(Spacer(1, 5 * mm))

        grid = _grid_map(instance, timetable, lambda s, p, b=batch: s.batch_id == b.id)
        data = [[""] + [DAYS[d] for d in range(cal.days)]]
        highlight: list[tuple[int, int]] = []
        # Parallel electives stack two subjects in one cell, so rows must grow
        # to fit their busiest cell or the text bleeds across the grid line.
        row_heights: list[float] = []

        for period in range(cal.periods):
            stacked = max(
                (len(grid.get((d, period), [])) for d in range(cal.days)), default=1
            )
            row_heights.append(max(15, 11 * max(stacked, 1) + 4) * mm)
            row: list = [
                Paragraph(
                    f"<b>{PERIOD_START[period]}</b><br/>{PERIOD_END[period]}", cellstyle
                )
            ]
            for d in range(cal.days):
                slot = cal.by_id[d * cal.periods + period]
                if slot.is_lunch:
                    row.append(Paragraph("<i>lunch</i>", cellstyle))
                    continue
                entries = grid.get((d, period), [])
                if not entries:
                    row.append("")
                    continue
                if any(s.id in changed for s, _, _ in entries):
                    highlight.append((d + 1, period + 1))
                parts = []
                for s, p, cont in entries:
                    tag = " (cont.)" if cont else ""
                    name = instance.faculty_by_id[s.faculty_id].name
                    parts.append(
                        f"<b>{s.subject_code}</b>{tag}<br/>{p.room_id}<br/>"
                        f"<font size=6>{name}</font>"
                    )
                row.append(Paragraph("<br/>".join(parts), cellstyle))
            data.append(row)

        col = (doc.width - 22 * mm) / cal.days
        table = Table(
            data,
            colWidths=[22 * mm] + [col] * cal.days,
            rowHeights=[8 * mm] + row_heights,
        )
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2c45")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 8),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#bfbfbf")),
            ("BACKGROUND", (0, 1), (0, -1), colors.HexColor("#f2f2f2")),
        ]
        for c, r in highlight:
            style.append(("BACKGROUND", (c, r), (c, r), colors.HexColor("#fff2cc")))
        table.setStyle(TableStyle(style))
        story.append(table)

    # --- verification page -------------------------------------------
    story.append(PageBreak())
    story.append(Paragraph("Verification", h1))
    story.append(
        Paragraph(
            "Recounted directly from the placements by a checker that shares no "
            "logic with the CP-SAT model.",
            sub,
        )
    )
    story.append(Spacer(1, 4 * mm))

    rows = [["Hard constraint", "Violations"]]
    rows += [[k.replace("_", " "), str(v)] for k, v in report.counts.items()]
    if diff:
        rows.append(["", ""])
        rows.append(["Sessions unchanged", f"{diff.unchanged} of {diff.total}"])
        rows.append(["Schedule retention", f"{diff.retention_pct:.1f}%"])
    check = Table(rows, colWidths=[80 * mm, 40 * mm])
    check.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2c45")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#bfbfbf")),
            ]
        )
    )
    story.append(check)

    doc.build(story)
    return buffer.getvalue()


# ----------------------------------------------------------------------
# Calendar (RFC 5545)
# ----------------------------------------------------------------------


def _escape(text: str) -> str:
    return (
        text.replace("\\", "\\\\")
        .replace(";", "\\;")
        .replace(",", "\\,")
        .replace("\n", "\\n")
    )


def to_ics(
    instance: Instance,
    timetable: Timetable,
    keep=None,
    term_start: date | None = None,
    weeks: int = 14,
    calendar_name: str = "ChronoSolve timetable",
) -> str:
    """Weekly-recurring calendar events, one per scheduled session.

    `term_start` anchors week 1; the Monday of that week is used so each session
    lands on its own weekday. Defaults to the Monday of the current week.
    """
    cal = instance.calendar
    keep = keep or (lambda s, p: True)
    anchor = term_start or date.today()
    monday = anchor - timedelta(days=anchor.weekday())
    stamp = datetime.now().strftime("%Y%m%dT%H%M%S")

    lines = [
        "BEGIN:VCALENDAR",
        "VERSION:2.0",
        "PRODID:-//ChronoSolve//Academic Scheduling//EN",
        "CALSCALE:GREGORIAN",
        f"X-WR-CALNAME:{_escape(calendar_name)}",
    ]

    for s in instance.sessions:
        p = timetable.placements.get(s.id)
        if p is None or not keep(s, p):
            continue
        slot = cal.by_id[p.timeslot_id]
        day = monday + timedelta(days=slot.day)
        start_h, start_m = PERIOD_START[slot.period].split(":")
        end_index = min(slot.period + s.duration - 1, cal.periods - 1)
        end_h, end_m = PERIOD_END[end_index].split(":")

        faculty = instance.faculty_by_id[s.faculty_id].name
        room = instance.room_by_id[p.room_id]

        lines += [
            "BEGIN:VEVENT",
            f"UID:{s.id}@chronosolve",
            f"DTSTAMP:{stamp}",
            f"DTSTART:{day:%Y%m%d}T{start_h}{start_m}00",
            f"DTEND:{day:%Y%m%d}T{end_h}{end_m}00",
            f"RRULE:FREQ=WEEKLY;COUNT={weeks}",
            f"SUMMARY:{_escape(f'{s.subject_code} {s.subject_name}')}",
            f"LOCATION:{_escape(f'{room.name} ({room.id})')}",
            f"DESCRIPTION:{_escape(f'{s.batch_id} - {faculty} - {s.duration}h')}",
            "END:VEVENT",
        ]

    lines.append("END:VCALENDAR")
    # RFC 5545 requires CRLF line endings.
    return "\r\n".join(lines) + "\r\n"
