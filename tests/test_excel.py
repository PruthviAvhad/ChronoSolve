"""Excel import and template export.

The round-trip test is the load-bearing one: the template is generated from a
live instance, so if export and import agree, the format a coordinator is handed
is exactly the format the importer accepts.
"""

from __future__ import annotations

import io
from datetime import time

import pytest
from openpyxl import Workbook, load_workbook

from backend.app.data.excel import read_workbook, to_workbook

HEADERS = {
    "Rooms": ["id", "name", "capacity", "type"],
    "Faculty": ["id", "name", "max_daily", "max_weekly", "max_consecutive"],
    "Batches": ["id", "name", "strength"],
    "Subjects": [
        "code",
        "name",
        "batch",
        "faculty",
        "sessions_per_week",
        "hours_per_session",
        "room_type",
        "elective_group",
        "headcount",
    ],
    "Unavailability": ["kind", "id", "day", "start", "end"],
}


def build(
    rooms=(("LH1", "Hall 1", 80, "LECTURE"), ("CL1", "Lab 1", 80, "LAB")),
    faculty=(("F1", "Prof. A", 5, 18, 3),),
    batches=(("B1", "Div 1", 60),),
    subjects=(("CS1", "Maths", "B1", "F1", 3, 1, "LECTURE", None, None),),
    unavailability=None,
    headers=None,
    omit=(),
) -> bytes:
    """Assemble a workbook from row tuples."""
    heads = {**HEADERS, **(headers or {})}
    wb = Workbook()
    wb.remove(wb.active)
    data = {
        "Rooms": rooms,
        "Faculty": faculty,
        "Batches": batches,
        "Subjects": subjects,
    }
    if unavailability is not None:
        data["Unavailability"] = unavailability

    for sheet, rowset in data.items():
        if sheet in omit:
            continue
        ws = wb.create_sheet(sheet)
        ws.append(heads[sheet])
        for row in rowset:
            ws.append(list(row))

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def messages(report) -> str:
    return " | ".join(i.message for i in report.issues)


# ----------------------------------------------------------------------
# round trip
# ----------------------------------------------------------------------


def test_the_template_reimports_to_an_identical_department(small_instance):
    book = to_workbook(small_instance)
    restored, report = read_workbook(book, name=small_instance.name)

    assert report.ok, report.render()
    assert restored is not None

    def signature(inst):
        return (
            sorted(
                (
                    s.id,
                    s.subject_code,
                    s.subject_name,
                    s.batch_id,
                    s.faculty_id,
                    s.duration,
                    s.room_type.value,
                    s.elective_group,
                    s.headcount,
                )
                for s in inst.sessions
            ),
            sorted(
                (r.id, r.name, r.capacity, r.room_type.value, sorted(r.unavailable))
                for r in inst.rooms
            ),
            sorted(
                (
                    f.id,
                    f.name,
                    sorted(f.unavailable),
                    f.max_daily_load,
                    f.max_weekly_load,
                    f.max_consecutive,
                )
                for f in inst.faculty
            ),
            sorted((b.id, b.name, b.strength) for b in inst.batches),
        )

    assert signature(restored) == signature(small_instance)


def test_the_template_carries_every_sheet_the_importer_needs(small_instance):
    wb = load_workbook(io.BytesIO(to_workbook(small_instance)))
    assert set(wb.sheetnames) == {
        "Rooms",
        "Faculty",
        "Batches",
        "Subjects",
        "Unavailability",
    }
    # Subjects collapse back to one row per weekly requirement, not per session.
    assert wb["Subjects"].max_row - 1 < len(small_instance.sessions)


# ----------------------------------------------------------------------
# expansion
# ----------------------------------------------------------------------


def test_a_subject_expands_into_its_weekly_sessions():
    inst, report = read_workbook(
        build(subjects=(("CS1", "Maths", "B1", "F1", 4, 1, "LECTURE", None, None),))
    )
    assert report.ok, report.render()
    assert len(inst.sessions) == 4
    assert {s.id for s in inst.sessions} == {f"B1-CS1-{k}" for k in range(1, 5)}


def test_a_multi_hour_lab_keeps_its_duration():
    inst, report = read_workbook(
        build(subjects=(("CS9", "Lab", "B1", "F1", 1, 2, "LAB", None, None),))
    )
    assert report.ok, report.render()
    assert inst.sessions[0].duration == 2
    assert inst.sessions[0].is_lab


def test_parallel_electives_are_grouped_per_week():
    inst, report = read_workbook(
        build(
            rooms=(("LH1", "Hall", 80, "LECTURE"), ("LH2", "Hall 2", 80, "LECTURE")),
            subjects=(
                ("CS4", "Cloud", "B1", "F1", 2, 1, "LECTURE", "ELEC1", 30),
                ("CS5", "Security", "B1", "F1", 2, 1, "LECTURE", "ELEC1", 30),
            ),
        )
    )
    assert report.ok, report.render()
    groups = inst.elective_groups()
    assert set(groups) == {"ELEC1-1", "ELEC1-2"}
    assert all(len(members) == 2 for members in groups.values())
    assert all(s.headcount == 30 for s in inst.sessions)


def test_unavailability_windows_become_blocked_slots():
    inst, report = read_workbook(
        build(unavailability=(("faculty", "F1", "Wed", "14:00", "17:00"),))
    )
    assert report.ok, report.render()
    blocked = inst.faculty_by_id["F1"].unavailable
    assert {inst.calendar.by_id[u].label for u in blocked} == {
        "Wed 14:00",
        "Wed 15:00",
        "Wed 16:00",
    }


@pytest.mark.parametrize(
    "start,end,expected",
    [
        ("2 PM", "5 PM", {"Wed 14:00", "Wed 15:00", "Wed 16:00"}),
        ("09:00", "11:00", {"Wed 09:00", "Wed 10:00"}),
        (time(14, 0), time(16, 0), {"Wed 14:00", "Wed 15:00"}),
    ],
)
def test_times_are_accepted_in_the_forms_people_type(start, end, expected):
    inst, report = read_workbook(
        build(unavailability=(("faculty", "F1", "Wed", start, end),))
    )
    assert report.ok, report.render()
    blocked = inst.faculty_by_id["F1"].unavailable
    assert {inst.calendar.by_id[u].label for u in blocked} == expected


def test_a_division_can_be_blocked_out_in_the_workbook():
    inst, report = read_workbook(
        build(unavailability=(("batch", "B1", "Wed", "14:00", "17:00"),))
    )
    assert report.ok, report.render()
    blocked = inst.batch_by_id["B1"].unavailable
    assert {inst.calendar.by_id[u].label for u in blocked} == {
        "Wed 14:00",
        "Wed 15:00",
        "Wed 16:00",
    }


def test_an_unknown_division_in_unavailability_is_rejected():
    inst, report = read_workbook(
        build(unavailability=(("batch", "GHOST", "Wed", "14:00", "17:00"),))
    )
    assert inst is None
    assert "unknown batch" in messages(report)


def test_column_aliases_are_accepted():
    headers = {
        "Rooms": ["id", "name", "seats", "room_type"],
        "Subjects": [
            "subject_code",
            "subject",
            "division",
            "teacher",
            "per_week",
            "hours",
            "room_type",
            "elective",
            "cohort",
        ],
    }
    inst, report = read_workbook(build(headers=headers))
    assert report.ok, report.render()
    assert len(inst.sessions) == 3


# ----------------------------------------------------------------------
# rejection
# ----------------------------------------------------------------------


def test_a_file_that_is_not_a_workbook_is_refused():
    inst, report = read_workbook(b"definitely not a spreadsheet")
    assert inst is None
    assert not report.ok
    assert "could not be opened" in messages(report)


def test_a_missing_sheet_names_what_is_missing():
    inst, report = read_workbook(build(omit=("Subjects",)))
    assert inst is None
    assert "Subjects" in messages(report)


def test_an_unknown_batch_is_reported_with_its_row():
    inst, report = read_workbook(
        build(subjects=(("CS1", "Maths", "GHOST", "F1", 3, 1, "LECTURE", None, None),))
    )
    assert inst is None
    assert report.errors[0].sheet == "Subjects"
    assert report.errors[0].row == 2
    assert "unknown batch" in report.errors[0].message


def test_an_unknown_faculty_is_reported():
    inst, report = read_workbook(
        build(subjects=(("CS1", "Maths", "B1", "F404", 3, 1, "LECTURE", None, None),))
    )
    assert inst is None
    assert "unknown faculty" in messages(report)


@pytest.mark.parametrize(
    "rooms,expected",
    [
        ((("LH1", "A", 80, "LECTURE"), ("LH1", "B", 70, "LECTURE")), "duplicate"),
        ((("LH1", "A", 80, "SPACESHIP"),), "LECTURE or LAB"),
        ((("LH1", "A", 0, "LECTURE"),), "positive"),
    ],
)
def test_bad_room_rows_are_rejected(rooms, expected):
    inst, report = read_workbook(build(rooms=rooms))
    assert inst is None
    assert expected in messages(report)


def test_a_subject_listed_twice_for_one_batch_is_rejected():
    inst, report = read_workbook(
        build(
            subjects=(
                ("CS1", "Maths", "B1", "F1", 3, 1, "LECTURE", None, None),
                ("CS1", "Maths again", "B1", "F1", 2, 1, "LECTURE", None, None),
            )
        )
    )
    assert inst is None
    assert "listed twice" in messages(report)


def test_a_session_nobody_can_host_is_rejected():
    inst, report = read_workbook(
        build(
            rooms=(("LH1", "Small", 20, "LECTURE"),),
            subjects=(("CS1", "Maths", "B1", "F1", 1, 1, "LECTURE", None, None),),
        )
    )
    assert inst is None
    assert "no lecture room seats 60" in messages(report)


def test_a_teacher_beyond_their_weekly_cap_is_rejected():
    inst, report = read_workbook(
        build(
            faculty=(("F1", "Prof. A", 5, 4, 3),),
            subjects=(("CS1", "Maths", "B1", "F1", 10, 1, "LECTURE", None, None),),
        )
    )
    assert inst is None
    assert "weekly cap" in messages(report)


def test_a_division_needing_more_than_a_week_is_rejected():
    inst, report = read_workbook(
        build(
            faculty=(("F1", "Prof. A", 8, 99, 8), ("F2", "Prof. B", 8, 99, 8)),
            subjects=(
                ("CS1", "Maths", "B1", "F1", 30, 1, "LECTURE", None, None),
                ("CS2", "Physics", "B1", "F2", 30, 1, "LECTURE", None, None),
            ),
        )
    )
    assert inst is None
    assert "only 35" in messages(report)


def test_every_error_is_collected_rather_than_only_the_first():
    inst, report = read_workbook(
        build(
            subjects=(
                ("CS1", "A", "GHOST", "F1", 1, 1, "LECTURE", None, None),
                ("CS2", "B", "B1", "F404", 1, 1, "LECTURE", None, None),
                ("CS3", "C", "B1", "F1", 1, 99, "LECTURE", None, None),
            )
        )
    )
    assert inst is None
    assert len(report.errors) >= 3


def test_a_warning_does_not_block_the_import():
    """An idle teacher is worth flagging but is not a reason to refuse."""
    inst, report = read_workbook(
        build(faculty=(("F1", "Prof. A", 5, 18, 3), ("F2", "Prof. Idle", 5, 18, 3)))
    )
    assert report.ok
    assert inst is not None
    assert any("teach nothing" in w.message for w in report.warnings)


def test_the_report_counts_what_it_read():
    _, report = read_workbook(build())
    assert report.counts == {"rooms": 2, "faculty": 1, "batches": 1, "sessions": 3}
    assert "3 sessions" in report.render()
