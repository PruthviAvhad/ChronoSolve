"""Explanations, diagnosis, language parsing, persistence and export."""

from __future__ import annotations

import io
import zipfile
from dataclasses import replace

import pytest

from backend.app.data.store import load_snapshot, save_snapshot
from backend.app.domain.models import Instance, RoomType
from backend.app.export import to_ics, to_pdf, to_xlsx
from backend.app.nl import parse
from backend.app.solver.diagnose import diagnose
from backend.app.solver.disruption import apply_disruptions, faculty_unavailable
from backend.app.solver.explain import explain_move, explain_session
from backend.app.solver.repair import repair


# ----------------------------------------------------------------------
# explanations
# ----------------------------------------------------------------------


def test_a_session_is_never_reported_as_blocking_itself(small_instance, published):
    """The probe lifts the session out before testing, so its own slot must
    come back available -- otherwise every explanation would be nonsense."""
    for session in small_instance.sessions[:6]:
        ex = explain_session(small_instance, published, session.id)
        here = next(
            o
            for o in ex.options
            if o.timeslot_id == published.placements[session.id].timeslot_id
        )
        assert here.feasible, f"{session.id} blocks itself: {here.blockers}"


def test_every_blocked_slot_states_a_reason(small_instance, published):
    ex = explain_session(small_instance, published, small_instance.sessions[0].id)
    assert len(ex.options) == len(small_instance.calendar)
    for option in ex.options:
        if not option.feasible:
            assert option.blockers, f"{option.label} blocked without a reason"
            assert all(b.rule and b.message for b in option.blockers)
        else:
            assert option.room_id is not None


def test_lunch_and_day_end_block_a_multi_hour_lab(small_instance, published):
    cal = small_instance.calendar
    lab = next(s for s in small_instance.sessions if s.duration > 1)
    ex = explain_session(small_instance, published, lab.id)
    by_slot = {o.timeslot_id: o for o in ex.options}

    lunch = next(sl for sl in cal.slots if sl.is_lunch)
    assert not by_slot[lunch.id].feasible

    last = next(sl for sl in cal.slots if sl.period == cal.periods - 1)
    assert "day_boundary" in {b.rule for b in by_slot[last.id].blockers}


def test_every_moved_session_can_explain_why(small_instance, published):
    victim = next(s for s in small_instance.sessions if s.duration == 1)
    slot = small_instance.calendar.by_id[published.placements[victim.id].timeslot_id]
    fac = small_instance.faculty_by_id[victim.faculty_id]
    disruptions = [
        faculty_unavailable(small_instance, fac.name, slot.day, "09:00", "23:59")
    ]
    disrupted = apply_disruptions(small_instance, disruptions)
    result = repair(disrupted, published, disruptions, phase1_limit=15, phase2_limit=5)
    assert result.solved

    for change in result.diff.changes:
        reasons = explain_move(
            disrupted,
            result.timetable,
            change.session_id,
            change.from_slot,
            change.from_room,
        )
        assert reasons, f"{change.session_id} moved with no explanation"
        assert all(r.message for r in reasons)


# ----------------------------------------------------------------------
# infeasibility diagnosis
# ----------------------------------------------------------------------


def test_a_healthy_instance_has_no_blocking_findings(small_instance):
    assert not diagnose(small_instance).proven_infeasible


def test_blocking_a_teacher_all_week_is_proven_impossible(small_instance):
    busiest = max(
        small_instance.faculty,
        key=lambda f: sum(
            s.duration for s in small_instance.sessions if s.faculty_id == f.id
        ),
    )
    disruptions = [
        faculty_unavailable(small_instance, busiest.name, day, "09:00", "23:59")
        for day in range(small_instance.calendar.days)
    ]
    report = diagnose(apply_disruptions(small_instance, disruptions))

    assert report.proven_infeasible
    assert "Faculty availability" in {f.category for f in report.blocking}
    # The most actionable finding is ranked first.
    assert report.blocking[0].category == "Faculty availability"
    assert all(f.suggestions for f in report.blocking)


def test_removing_every_laboratory_is_proven_impossible(small_instance):
    """Labs cannot run in lecture rooms, so deleting the labs must be caught."""
    stripped = Instance(
        name=small_instance.name,
        calendar=small_instance.calendar,
        rooms=[r for r in small_instance.rooms if r.room_type is not RoomType.LAB],
        faculty=list(small_instance.faculty),
        batches=list(small_instance.batches),
        sessions=list(small_instance.sessions),
    )
    report = diagnose(stripped)
    assert report.proven_infeasible
    assert any(
        "room" in f.category.lower() or "placement" in f.category.lower()
        for f in report.blocking
    )


def test_an_overloaded_teacher_is_caught_by_the_weekly_cap(small_instance):
    victim = small_instance.faculty[0]
    tightened = Instance(
        name=small_instance.name,
        calendar=small_instance.calendar,
        rooms=list(small_instance.rooms),
        faculty=[
            replace(f, max_weekly_load=1) if f.id == victim.id else f
            for f in small_instance.faculty
        ],
        batches=list(small_instance.batches),
        sessions=list(small_instance.sessions),
    )
    report = diagnose(tightened)
    assert report.proven_infeasible
    assert any(f.category == "Weekly teaching load" for f in report.blocking)


# ----------------------------------------------------------------------
# natural-language parsing
# ----------------------------------------------------------------------


@pytest.mark.parametrize(
    "sentence,kind,days,start,end",
    [
        (
            "Prof. Mehta is unavailable after 2 PM on Friday",
            "FACULTY_UNAVAILABLE",
            [4],
            "14:00",
            "23:59",
        ),
        (
            "Prof. Mehta is unavailable on Wednesday from 11:00 to 14:00",
            "FACULTY_UNAVAILABLE",
            [2],
            "11:00",
            "14:00",
        ),
        (
            "Mehta is on leave on Monday and Tuesday",
            "FACULTY_UNAVAILABLE",
            [0, 1],
            "09:00",
            "23:59",
        ),
        (
            "SE-A is at a seminar on Wednesday afternoon",
            "BATCH_UNAVAILABLE",
            [2],
            "13:00",
            "23:59",
        ),
        (
            "Computer Lab 3 is closed on Friday",
            "ROOM_UNAVAILABLE",
            [4],
            "09:00",
            "23:59",
        ),
        (
            "CL3 is under maintenance on Wednesday morning",
            "ROOM_UNAVAILABLE",
            [2],
            "09:00",
            "13:00",
        ),
        (
            "Lecture Hall 108 is unavailable Monday before 11 AM",
            "ROOM_UNAVAILABLE",
            [0],
            "09:00",
            "11:00",
        ),
    ],
)
def test_sentences_become_the_expected_rule(
    full_instance, sentence, kind, days, start, end
):
    rule = parse(sentence, full_instance)
    assert rule.understood, rule.issues
    assert rule.kind == kind
    assert rule.days == days
    assert (rule.start_time, rule.end_time) == (start, end)


def test_one_sentence_can_yield_several_rules(full_instance):
    rules = parse("Mehta is on leave on Monday and Tuesday", full_instance)
    emitted = rules.to_disruptions()
    assert len(emitted) == 2
    assert {r["day"] for r in emitted} == {0, 1}
    assert all(
        set(r) == {"kind", "target", "day", "start_time", "end_time"} for r in emitted
    )


@pytest.mark.parametrize(
    "sentence,expected",
    [
        ("Prof. Nobody is unavailable on Friday", "recognised"),
        ("Prof. Mehta is unavailable", "weekday"),
        ("Prof. Mehta teaches on Friday", "unavailability"),
        ("", "sentence"),
    ],
)
def test_unparseable_sentences_explain_what_is_missing(
    full_instance, sentence, expected
):
    rule = parse(sentence, full_instance)
    assert not rule.understood
    assert rule.to_disruptions() == []
    assert any(expected in issue for issue in rule.issues), rule.issues


def test_an_inferred_meridiem_is_disclosed(full_instance):
    rule = parse("Prof. Mehta is unavailable after 2 on Friday", full_instance)
    assert rule.understood
    assert rule.start_time == "14:00"
    assert rule.assumptions, "inferring 2 -> 14:00 must be reported"


# ----------------------------------------------------------------------
# persistence
# ----------------------------------------------------------------------


def test_snapshot_round_trip_is_lossless(small_instance, published, tmp_path):
    path = save_snapshot(
        tmp_path / "snap.json", small_instance, published, label="test-v1"
    )
    loaded, timetable, meta = load_snapshot(path)

    assert meta["label"] == "test-v1"
    assert meta["saved_at"].endswith("Z")
    assert len(loaded.sessions) == len(small_instance.sessions)
    assert {r.id for r in loaded.rooms} == {r.id for r in small_instance.rooms}
    assert timetable.placements == published.placements
    assert timetable.status == published.status
    for a, b in zip(
        sorted(loaded.faculty, key=lambda f: f.id),
        sorted(small_instance.faculty, key=lambda f: f.id),
    ):
        assert a.unavailable == b.unavailable
        assert a.max_weekly_load == b.max_weekly_load


def test_an_unknown_schema_version_is_refused(small_instance, published, tmp_path):
    path = save_snapshot(tmp_path / "snap.json", small_instance, published)
    path.write_text(
        path.read_text(encoding="utf-8").replace(
            '"schema_version": 1', '"schema_version": 99'
        ),
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="schema"):
        load_snapshot(path)


# ----------------------------------------------------------------------
# export
# ----------------------------------------------------------------------


def test_excel_export_has_a_sheet_per_division(small_instance, published):
    payload = to_xlsx(small_instance, published, label="test")
    with zipfile.ZipFile(io.BytesIO(payload)) as zf:
        assert "xl/workbook.xml" in zf.namelist()

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(payload))
    assert "Summary" in wb.sheetnames
    assert "All sessions" in wb.sheetnames
    for b in small_instance.batches:
        assert b.id in wb.sheetnames
    # One row per placed session, plus the header.
    assert wb["All sessions"].max_row == len(published.placements) + 1


def test_pdf_export_is_a_valid_document(small_instance, published):
    payload = to_pdf(small_instance, published, label="test")
    assert payload.startswith(b"%PDF")
    assert payload.rstrip().endswith(b"%%EOF")
    assert payload.count(b"/Type /Page") >= len(small_instance.batches)


def test_calendar_export_is_well_formed(small_instance, published):
    text = to_ics(small_instance, published)
    assert text.startswith("BEGIN:VCALENDAR")
    assert text.rstrip().endswith("END:VCALENDAR")
    assert text.count("BEGIN:VEVENT") == len(published.placements)
    assert text.count("BEGIN:VEVENT") == text.count("END:VEVENT")
    assert "RRULE:FREQ=WEEKLY" in text
    # RFC 5545 requires CRLF line endings.
    assert "\r\n" in text
    assert "\n" not in text.replace("\r\n", "")


def test_calendar_export_can_be_filtered_to_one_person(small_instance, published):
    target = small_instance.sessions[0].faculty_id
    expected = sum(1 for s in small_instance.sessions if s.faculty_id == target)
    text = to_ics(small_instance, published, keep=lambda s, p: s.faculty_id == target)
    assert text.count("BEGIN:VEVENT") == expected
